from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.core.paginator import Paginator
from django.db import transaction
from django.shortcuts import get_object_or_404, redirect, render
from django.views.decorators.http import require_POST
from openpyxl import load_workbook

from .forms import ExcelUploadForm, StudentForm, StudentSearchForm
from .models import Student


def home(request):
    form = StudentSearchForm(request.POST or None)
    student = None
    searched = False
    if request.method == 'POST' and form.is_valid():
        searched = True
        admission_no = form.cleaned_data['admission_no'].strip()
        student = Student.objects.filter(admission_no__iexact=admission_no).first()
    student_fields = []
    if student:
        student_fields = [('Admission No', student.admission_no), ('Full Name', student.full_name), ('Course Code', student.course_code), ('Campus', student.campus), ('Supervisor', student.supervisor)]
    return render(request, 'students/home.html', {'form': form, 'student': student, 'searched': searched, 'student_fields': student_fields})

@login_required
def dashboard(request):
    q = request.GET.get('q', '').strip()
    students = Student.objects.all()
    if q:
        students = students.filter(admission_no__icontains=q) | students.filter(full_name__icontains=q)
    paginator = Paginator(students.distinct(), 15)
    page = paginator.get_page(request.GET.get('page'))
    return render(request, 'students/dashboard.html', {'page': page, 'q': q, 'count': Student.objects.count()})

@login_required
def student_create(request):
    form = StudentForm(request.POST or None)
    if form.is_valid():
        form.save(); messages.success(request, 'Student added successfully.'); return redirect('dashboard')
    return render(request, 'students/student_form.html', {'form': form, 'title': 'Add Student'})

@login_required
def student_edit(request, pk):
    student = get_object_or_404(Student, pk=pk)
    form = StudentForm(request.POST or None, instance=student)
    if form.is_valid():
        form.save(); messages.success(request, 'Student updated successfully.'); return redirect('dashboard')
    return render(request, 'students/student_form.html', {'form': form, 'title': 'Edit Student', 'student': student})

@login_required
@require_POST
def student_delete(request, pk):
    student = get_object_or_404(Student, pk=pk)
    student.delete(); messages.success(request, 'Student deleted successfully.')
    return redirect('dashboard')

@login_required
def import_excel(request):
    form = ExcelUploadForm(request.POST or None, request.FILES or None)
    result = None
    if request.method == 'POST' and form.is_valid():
        uploaded = form.cleaned_data['file']
        if not uploaded.name.lower().endswith('.xlsx'):
            form.add_error('file', 'Only .xlsx files are supported.')
        else:
            try:
                wb = load_workbook(uploaded, read_only=True, data_only=True)
                ws = wb.active
                rows = ws.iter_rows(values_only=True)
                headers = [str(x).strip().lower() if x is not None else '' for x in next(rows)]
                required = ['admission no', 'fullname', 'course code', 'campus', 'supervisor']
                # Accept Full Name as an alternative header spelling.
                aliases = {'full name': 'fullname', 'full_name': 'fullname', 'admission number': 'admission no', 'course': 'course code'}
                headers = [aliases.get(h, h) for h in headers]
                missing = [h for h in required if h not in headers]
                if missing:
                    form.add_error('file', 'Missing columns: ' + ', '.join(missing))
                else:
                    idx = {h: headers.index(h) for h in required}
                    created = updated = skipped = 0
                    with transaction.atomic():
                        for row in rows:
                            values = [row[idx[h]] if idx[h] < len(row) else None for h in required]
                            values = [str(v).strip() if v is not None else '' for v in values]
                            if not values[0] or not values[1]:
                                skipped += 1; continue
                            defaults = {'full_name': values[1], 'course_code': values[2], 'campus': values[3], 'supervisor': values[4]}
                            obj = Student.objects.filter(admission_no__iexact=values[0]).first()
                            if obj and form.cleaned_data['replace_existing']:
                                for k, v in defaults.items(): setattr(obj, k, v)
                                obj.save(); updated += 1
                            elif obj:
                                skipped += 1
                            else:
                                Student.objects.create(admission_no=values[0], **defaults); created += 1
                    result = {'created': created, 'updated': updated, 'skipped': skipped}
                    messages.success(request, f'Import complete: {created} added, {updated} updated, {skipped} skipped.')
            except Exception as exc:
                form.add_error('file', f'Could not read the workbook: {exc}')
    return render(request, 'students/import.html', {'form': form, 'result': result})
